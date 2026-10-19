"""把 textbooks/*.json 統整成 1019課程版本.xlsx

用法：python3 data/build_textbooks.py
含公式的「涵蓋總覽」需用 LibreOffice 重算（或在 Excel 開啟時自動計算）。
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
SRC = BASE / "textbooks"
OUT = BASE / "1019課程版本.xlsx"

COLS = ["版本", "學段", "年級", "冊次", "科目", "章", "節", "學年度", "來源類型", "來源網址", "備註"]
WIDTHS = [8, 6, 6, 18, 10, 34, 60, 8, 14, 50, 24]
PUB_ORDER = ["翰林", "康軒", "南一", "三民", "龍騰", "全華", "泰宇"]
SUBJ_ORDER = ["國文", "英文", "數學", "生物", "理化", "物理", "化學", "地球科學", "地科",
              "歷史", "地理", "公民", "公民與社會"]
SUBJ_ALIAS = {"自然-生物": "生物", "自然-理化": "理化", "自然-地球科學": "地球科學", "地科": "地球科學"}
VOL_ORDER = {"上": 0, "下": 1}

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
EMPTY_FILL = PatternFill("solid", start_color="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def load_rows():
    rows, seen = [], set()
    for f in sorted(SRC.glob("*.json")):
        data = json.loads(f.read_text(encoding="utf-8"))
        for i, r in enumerate(data):
            row = {c: norm(r.get(c)) for c in COLS}
            row["科目"] = SUBJ_ALIAS.get(row["科目"], row["科目"])
            if not (row["版本"] and row["學段"] and row["科目"]):
                print(f"skip (missing key fields) {f.name}#{i}", file=sys.stderr)
                continue
            key = tuple(row[c] for c in ["版本", "學段", "年級", "冊次", "科目", "章", "節"])
            if key in seen:
                continue
            seen.add(key)
            row["_order"] = (f.name, i)
            rows.append(row)
    # 同一冊若已有章節，就移除該冊的「未查到」空白列
    found = {(r["版本"], r["學段"], r["年級"], r["冊次"], r["科目"]) for r in rows if r["章"]}
    rows = [r for r in rows
            if r["章"] or (r["版本"], r["學段"], r["年級"], r["冊次"], r["科目"]) not in found]
    return rows


def idx(lst, v):
    return lst.index(v) if v in lst else len(lst)


def sort_key(r):
    grade = int(r["年級"]) if r["年級"] and r["年級"].isdigit() else 99
    return (idx(PUB_ORDER, r["版本"]), idx(SUBJ_ORDER, r["科目"]), r["科目"], grade,
            VOL_ORDER.get(r["冊次"] or "", 5), r["冊次"] or "", r["_order"])


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"


def write_data_sheet(ws, rows):
    ws.append(COLS)
    style_header(ws, len(COLS))
    for r in rows:
        ws.append([r[c] for c in COLS])
        n = ws.max_row
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=n, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if not r["章"]:
                cell.fill = EMPTY_FILL
        url = r["來源網址"]
        if url and url.startswith("http"):
            link = ws.cell(row=n, column=10)
            link.hyperlink = url
            link.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"


def write_coverage(ws, rows):
    heads = ["學段", "版本", "科目", "年級", "冊次", "已查到章（課）數", "狀態"]
    ws.append(heads)
    style_header(ws, len(heads))
    books, seen = [], set()
    for r in rows:
        k = (r["學段"], r["版本"], r["科目"], r["年級"], r["冊次"])
        if k not in seen:
            seen.add(k)
            books.append(k)
    for stage, pub, subj, grade, vol in books:
        sheet = "國中" if stage == "國中" else "高中"
        n = ws.max_row + 1
        # 用 COUNTIFS 直接數資料表，資料表增修時會自動更新
        crit = [f'{sheet}!$A:$A,$B{n}', f'{sheet}!$E:$E,$C{n}', f'{sheet}!$F:$F,"<>"']
        crit.append(f'{sheet}!$C:$C,$D{n}' if grade else f'{sheet}!$C:$C,""')
        crit.append(f'{sheet}!$D:$D,$E{n}' if vol else f'{sheet}!$D:$D,""')
        ws.append([stage, pub, subj, grade, vol, f"=COUNTIFS({','.join(crit)})",
                   f'=IF(F{n}>0,"已查到","未查到")'])
        for c in range(1, len(heads) + 1):
            cell = ws.cell(row=n, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
    for i, w in enumerate([6, 8, 10, 6, 18, 16, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:G{ws.max_row}"


def write_notes(ws, rows):
    found = sum(1 for r in rows if r["章"])
    lines = [
        ("1019課程版本：各出版商國高中課本目錄統整", True),
        ("", False),
        ("資料範圍", True),
        ("國中：翰林、康軒、南一，7–9 年級上下冊。國文、英文、數學、自然（生物、理化、地球科學）、社會（歷史、地理、公民）。", False),
        ("高中：三民、龍騰、南一、全華、泰宇。10 科部定必修與加深加廣選修。", False),
        ("以 115 學年度現行版本為準；只找到 114 學年度資料時，「學年度」欄會註明。", False),
        ("", False),
        ("資料來源與限制", True),
        ("蒐集時，出版商官網在作業環境中無法連線（網路政策阻擋），因此本版資料來自網路搜尋，"
         "主要是各校公開的課程計畫與教學計畫表，每列都附來源網址。", False),
        ("只收錄來源中實際看到、且明確標示版本的章節；查不到的冊次保留一列，章節空白（黃底），不做推測。", False),
        ("學校課程計畫的單元名稱可能經過學校簡化或調整，與課本原文可能有些微差異；以出版商官網或課本為準。", False),
        ("待辦：開放出版商網域後，到官網逐冊核對，並補上空白冊次；核對過的列，「來源類型」會改為「出版商官網」。", False),
        ("官網網域：翰林 hle.com.tw、康軒 knsh.com.tw、南一 nani.com.tw、三民 sanmin.com.tw、龍騰 ltedu.com.tw、全華 chwa.com.tw、泰宇 tyedu.com.tw（泰宇待確認）。", False),
        ("", False),
        ("欄位說明", True),
        ("「國中」「高中」：一章（國文、英文為一課）一列；「節」以全形分號「；」分隔。", False),
        ("「涵蓋總覽」：每冊已查到的章數，以公式從資料表計算，資料增修後會自動更新。", False),
        ("", False),
        (f"目前共 {len(rows)} 列，其中 {found} 列有章節內容。", False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=bold, size=12 if bold else 10)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110


def main():
    rows = sorted(load_rows(), key=sort_key)
    wb = Workbook()
    notes = wb.active
    notes.title = "說明"
    write_notes(notes, rows)
    write_data_sheet(wb.create_sheet("國中"), [r for r in rows if r["學段"] == "國中"])
    write_data_sheet(wb.create_sheet("高中"), [r for r in rows if r["學段"] != "國中"])
    write_coverage(wb.create_sheet("涵蓋總覽"), rows)
    wb.save(OUT)
    print(f"wrote {OUT} rows={len(rows)} with_chapters={sum(1 for r in rows if r['章'])}")


if __name__ == "__main__":
    main()
